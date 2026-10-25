"""Trusted read and draft tools exposed to Kloppy's Strands agent."""

import json
import logging
import re
from base64 import b64encode
from contextlib import ExitStack
from io import BytesIO
from typing import Any, Literal
from urllib.parse import urlparse

import requests
from docx import Document
from docx.document import Document as DocxDocument
from strands import ToolContext, tool
from strands.tools.mcp import MCPClient
from tavily import TavilyClient
from tavily.errors import (
    BadRequestError,
    ForbiddenError,
    InvalidAPIKeyError,
    MissingAPIKeyError,
    TavilyKeylessLimitError,
    TimeoutError as TavilyTimeoutError,
    UsageLimitExceededError,
)

from ..config import settings
from ..observability import log_external_call
from .mcp_connections import user_server_config
from .memory import remember_project_fact
from .qloo import (
    ENTITY_TYPES,
    QlooConfigurationError,
    QlooProviderError,
    agent_grounding_options,
    run_capability,
    search_entities,
)
from ..schemas.qloo import QlooCapabilityRequest

logger = logging.getLogger(__name__)
CAMPAIGN_SECTIONS = (
    ("overview", "Overview"),
    ("audience", "Audience"),
    ("strategy", "Strategy"),
    ("launch", "Launch & Learn"),
)


@tool
def qloo_search_entities(query: str, entity_type: str = "urn:entity:brand") -> dict[str, Any]:
    """Find Qloo entities by name to ground an audience insight in real cultural entities."""
    if not settings.qloo_api_key:
        return {"available": False, "message": "Qloo is not configured."}
    try:
        return {
            "source": "Qloo API entity search",
            "entities": search_entities(
                query,
                next((name for name, urn in ENTITY_TYPES.items() if urn == entity_type), "brand"),
                take=8,
            ),
        }
    except (QlooProviderError, ValueError) as error:
        logger.warning("Qloo entity search failed: %s", type(error).__name__)
        return {"source": "Qloo API entity search", "error": "Qloo search is temporarily unavailable."}


@tool
def qloo_cultural_affinities(
    entity_ids: list[str],
    result_type: str = "urn:entity:brand",
    location_query: str = "",
    take: int = 10,
) -> dict[str, Any]:
    """Get Qloo cultural affinity recommendations from selected entities and an optional location."""
    if not settings.qloo_api_key:
        return {"available": False, "message": "Qloo is not configured."}
    safe_ids = [value[:160] for value in entity_ids[:8] if isinstance(value, str) and value.strip()]
    if not safe_ids:
        return {"error": "Provide one or more Qloo entity IDs from entity search."}
    try:
        category = next(
            (name for name, urn in ENTITY_TYPES.items() if urn == result_type),
            "brand",
        )
        response = run_capability(
            QlooCapabilityRequest(
                capability="locality" if location_query.strip() else "recommendations",
                entity_ids=safe_ids,
                result_category=category,
                location=location_query[:160] if location_query.strip() else None,
                take=max(1, min(take, 20)),
            )
        )
        return {
            "source": response["source"],
            "interpretation": response["interpretation"],
            "entities": response["results"],
            **({"heatmap": response["heatmap"]} if "heatmap" in response else {}),
        }
    except (QlooConfigurationError, QlooProviderError, ValueError) as error:
        logger.warning("Qloo affinity request failed: %s", type(error).__name__)
        return {"source": "Qloo API cultural affinity insights", "error": "Qloo insights are temporarily unavailable."}


@tool
def qloo_agent_grounding(query: str, entity_category: str = "brand", take: int = 8) -> dict[str, Any]:
    """Return a short, Qloo-ranked list of factual cultural options for campaign research."""
    if not settings.qloo_api_key:
        return {"available": False, "message": "Qloo is not configured."}
    try:
        return agent_grounding_options(query, entity_category, take)
    except (QlooConfigurationError, QlooProviderError, ValueError) as error:
        logger.warning("Qloo agent grounding failed: %s", type(error).__name__)
        return {"source": "Qloo API", "error": "Qloo grounding is temporarily unavailable."}


def _add_docx_body(document: DocxDocument, content: str) -> None:
    for line in content.splitlines():
        text = line.strip()
        if not text:
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", text)
        if heading:
            document.add_heading(heading.group(2).strip(), level=min(len(heading.group(1)), 3))
            continue
        bullet = re.match(r"^(?:[-*+])\s+(.+)$", text)
        if bullet:
            document.add_paragraph(bullet.group(1), style="List Bullet")
            continue
        numbered = re.match(r"^\d+[.)]\s+(.+)$", text)
        if numbered:
            document.add_paragraph(numbered.group(1), style="List Number")
            continue
        document.add_paragraph(text)


def _build_docx_artifact(title: str, sections: dict[str, str]) -> dict[str, Any]:
    safe_title = " ".join(title.split())[:140] or "Kloppy draft"
    filename = re.sub(r"[^a-z0-9]+", "-", safe_title.lower()).strip("-")[:90] or "kloppy-draft"
    normalized_sections = {
        key: str(sections.get(key) or "").strip()[:12000]
        for key, _ in CAMPAIGN_SECTIONS
    }
    document = Document()
    document.add_heading(safe_title, level=0)
    for key, heading in CAMPAIGN_SECTIONS:
        document.add_heading(heading, level=1)
        _add_docx_body(document, normalized_sections[key] or "Not developed in this draft.")
    output = BytesIO()
    document.save(output)
    content = "\n\n".join(
        f"## {heading}\n\n{normalized_sections[key] or 'Not developed in this draft.'}"
        for key, heading in CAMPAIGN_SECTIONS
    )
    return {
        "title": safe_title,
        "content": content,
        "format": "docx",
        "filename": f"{filename}.docx",
        "docx_base64": b64encode(output.getvalue()).decode("ascii"),
        "sections": normalized_sections,
    }


@tool(context=True)
def draft_document(
    title: str,
    overview: str,
    audience: str,
    strategy: str,
    launch: str,
    tool_context: ToolContext,
) -> dict[str, str]:
    """Create a reviewable Word document containing Overview, Audience, Strategy, and Launch & Learn."""
    artifact = _build_docx_artifact(
        title,
        {
            "overview": overview,
            "audience": audience,
            "strategy": strategy,
            "launch": launch,
        },
    )
    normalized_sections = artifact["sections"]
    state = tool_context.invocation_state
    state.setdefault("artifacts", []).append(artifact)
    state.setdefault("campaign_sections", {}).update(normalized_sections)
    return {"status": "drafted", "title": artifact["title"], "format": "Word document"}


def _tavily_search(query: str, operation: str, **options: Any) -> dict[str, Any] | None:
    try:
        with log_external_call("tavily", operation):
            response = TavilyClient(api_key=settings.tavily_api_key).search(
                query,
                timeout=15,
                **options,
            )
    except (
        BadRequestError,
        ForbiddenError,
        InvalidAPIKeyError,
        MissingAPIKeyError,
        TavilyKeylessLimitError,
        TavilyTimeoutError,
        UsageLimitExceededError,
        requests.RequestException,
        ValueError,
    ) as error:
        logger.warning("Tavily %s request failed: %s", operation, type(error).__name__)
        return None
    if not isinstance(response, dict):
        logger.warning("Tavily %s returned an invalid response shape", operation)
        return None
    return response


def _safe_web_url(value: Any) -> str | None:
    if not isinstance(value, str) or len(value) > 2000:
        return None
    parsed = urlparse(value)
    if (
        parsed.scheme not in {"http", "https"}
        or not parsed.hostname
        or parsed.username
        or parsed.password
    ):
        return None
    return value


def _run_web_search(query: str) -> dict[str, Any]:
    safe_query = " ".join(query.split())[:500]
    if not safe_query:
        return {"status": "error", "message": "Enter a search query."}
    if not settings.tavily_api_key:
        logger.warning("Web search unavailable: TAVILY_API_KEY is not configured")
        return {"status": "unavailable", "message": "Web search is not configured."}

    payload = _tavily_search(
        safe_query,
        "search",
        search_depth="basic",
        max_results=6,
        include_answer=False,
        include_raw_content=False,
        include_published_date=True,
    )
    if payload is None:
        return {"status": "error", "message": "Web search is temporarily unavailable."}
    results = payload.get("results", [])
    if not isinstance(results, list):
        logger.warning("Tavily search returned an invalid results shape")
        return {"status": "error", "message": "Web search returned an invalid response."}
    sources = [
        {
            "title": str(item.get("title") or "Untitled source")[:300],
            "url": _safe_web_url(item.get("url")) or "",
            "snippet": str(item.get("content") or "")[:2500],
            "published_date": str(item.get("published_date") or "")[:80],
        }
        for item in results[:6]
        if isinstance(item, dict) and _safe_web_url(item.get("url"))
    ]
    return {"status": "ok", "query": safe_query, "sources": sources}


@tool
def web_search(query: str) -> dict[str, Any]:
    """Search the web only when current, external, or explicitly requested facts would improve the answer."""
    return _run_web_search(query)


def _run_web_image_search(query: str) -> dict[str, Any]:
    safe_query = " ".join(query.split())[:500]
    if not safe_query:
        return {"status": "error", "message": "Enter an image search query."}
    if not settings.tavily_api_key:
        logger.warning("Image search unavailable: TAVILY_API_KEY is not configured")
        return {"status": "unavailable", "message": "Web image search is not configured."}

    payload = _tavily_search(
        safe_query,
        "image_search",
        search_depth="basic",
        max_results=6,
        include_answer=False,
        include_raw_content=False,
        include_images=True,
        include_image_descriptions=True,
        safe_search=True,
    )
    if payload is None:
        return {"status": "error", "message": "Web image search is temporarily unavailable."}

    source_by_image: dict[str, tuple[str, str]] = {}
    results = payload.get("results", [])
    if isinstance(results, list):
        for result in results:
            if not isinstance(result, dict):
                continue
            source_url = _safe_web_url(result.get("url"))
            source_title = str(result.get("title") or "Web image source")[:300]
            result_images = result.get("images", [])
            if not isinstance(result_images, list):
                continue
            for image in result_images:
                image_url = _safe_web_url(
                    image.get("url") if isinstance(image, dict) else image
                )
                if image_url and source_url:
                    source_by_image[image_url] = (source_title, source_url)

    images: list[dict[str, str]] = []
    seen: set[str] = set()
    image_results = payload.get("images", [])
    if isinstance(image_results, list):
        for image in image_results:
            image_url = _safe_web_url(
                image.get("url") if isinstance(image, dict) else image
            )
            if not image_url or image_url in seen:
                continue
            seen.add(image_url)
            source_title, source_url = source_by_image.get(
                image_url,
                ("Open original image", image_url),
            )
            description = (
                image.get("description", "")
                if isinstance(image, dict) and isinstance(image.get("description"), str)
                else ""
            )
            images.append(
                {
                    "url": image_url,
                    "description": description[:500] or safe_query,
                    "source_title": source_title,
                    "source_url": source_url,
                }
            )
            if len(images) == 6:
                break
    return {"status": "ok", "query": safe_query, "images": images}


@tool(context=True)
def web_image_search(query: str, tool_context: ToolContext) -> dict[str, Any]:
    """Find a small set of relevant web images to show as visual examples in chat."""
    result = _run_web_image_search(query)
    if result.get("status") == "ok":
        visual_results = tool_context.invocation_state.setdefault("visual_results", [])
        visual_results.extend(result["images"][: max(0, 6 - len(visual_results))])
    return result


@tool(context=True)
def walrus_remember_project_fact(fact: str, tool_context: ToolContext) -> dict[str, str]:
    """Remember a concise, explicit user-provided fact for the currently selected project."""
    namespace = tool_context.invocation_state.get("walrus_project_namespace")
    return remember_project_fact(
        namespace if isinstance(namespace, str) else None,
        fact,
    )


@tool(context=True)
def ask_user(
    question: str,
    question_type: Literal["clarification", "permission"],
    options: list[str],
    tool_context: ToolContext,
) -> dict[str, Any]:
    """Pause and ask the user a blocking clarification or permission question."""
    question = question.strip()
    if not question or len(question) > 1000:
        return {"status": "invalid_question"}
    if len(options) > 5 or any(not option.strip() or len(option) > 120 for option in options):
        return {"status": "invalid_options"}

    request_user_input = tool_context.invocation_state.get("request_user_input")
    if not callable(request_user_input):
        return {
            "status": "unavailable",
            "message": "Interactive questions require an active chat session.",
        }

    if question_type == "permission":
        options = ["Allow once", "Do not allow"]
    return request_user_input(
        {
            "question": question,
            "question_type": question_type,
            "options": options,
        }
    )


def _load_mcp_config(access_token: str | None = None) -> dict[str, Any]:
    raw: Any = {}
    if settings.mcp_servers_json:
        try:
            raw = json.loads(settings.mcp_servers_json)
        except json.JSONDecodeError as error:
            logger.error("MCP_SERVERS_JSON is invalid JSON at character %s", error.pos)
    servers = raw.get("mcpServers", raw) if isinstance(raw, dict) else {}
    if not isinstance(servers, dict):
        logger.error("MCP_SERVERS_JSON must contain an mcpServers object")
        servers = {}
    if access_token:
        servers = {**servers, **user_server_config(access_token).get("mcpServers", {})}
    safe_servers: dict[str, Any] = {}
    for name, original in servers.items():
        if not isinstance(original, dict):
            continue
        server = dict(original)
        server_url = server.get("url")
        parsed_url = urlparse(server_url) if isinstance(server_url, str) else None
        local_http = bool(parsed_url and parsed_url.scheme == "http" and parsed_url.hostname in {"localhost", "127.0.0.1", "::1"})
        if "command" in server or not parsed_url or (parsed_url.scheme != "https" and not local_http):
            logger.warning("MCP server %s skipped: configure a remote HTTPS or local loopback URL", name)
            continue
        allowed = server.pop("allowed_tools", [])
        if not isinstance(allowed, list) or not allowed:
            logger.info("MCP server %s skipped because allowed_tools is empty", name)
            continue
        prefix = re.sub(r"[^A-Za-z0-9_-]", "_", str(name))
        server["prefix"] = prefix
        server["tool_filters"] = {"allowed": [f"{prefix}_{item}" for item in allowed]}
        server["continue_on_error"] = True
        safe_servers[str(name)] = server
    return {"mcpServers": safe_servers}


def connected_mcp_clients(access_token: str | None = None) -> list[MCPClient]:
    """Create clients only for explicitly configured, allowlisted MCP tools."""
    config = _load_mcp_config(access_token)
    if not config["mcpServers"]:
        return []
    try:
        with log_external_call("mcp", "load_configured_servers"):
            return MCPClient.load_servers(config, continue_on_error=True, prefix_with_server_name=True)
    except (ValueError, TypeError) as error:
        logger.error("MCP server configuration could not be loaded: %s", type(error).__name__)
        return []


def agent_tool_session(
    access_token: str | None = None,
    memory_enabled: bool = False,
    interactive: bool = False,
) -> tuple[ExitStack, list[Any]]:
    """Open configured research tools and any available interactive tools."""
    stack = ExitStack()
    tools: list[Any] = [draft_document]
    if interactive:
        tools.append(ask_user)
    if memory_enabled:
        tools.append(walrus_remember_project_fact)
    tools.append(web_search)
    tools.append(web_image_search)
    if settings.qloo_api_key:
        tools.extend([qloo_search_entities, qloo_cultural_affinities, qloo_agent_grounding])
    clients = connected_mcp_clients(access_token)
    for client in clients:
        try:
            with log_external_call("mcp", "discover_read_only_tools"):
                stack.enter_context(client)
                for server_tool in client.list_tools_sync():
                    annotations = getattr(server_tool.mcp_tool, "annotations", None)
                    if annotations and getattr(
                        annotations,
                        "read_only_hint",
                        getattr(annotations, "readOnlyHint", False),
                    ):
                        tools.append(server_tool)
        except Exception as error:  # Keep healthy providers available when one MCP is offline.
            logger.warning("MCP connection unavailable: %s", type(error).__name__)
    return stack, tools
